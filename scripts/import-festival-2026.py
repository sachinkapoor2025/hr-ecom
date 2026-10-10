"""One-shot prod import for the 2026 Bhai Dooj / Diwali Orange County catalog.

Mirrors scripts/import-festival-2026.ts. Writes scripts/data/festival-2026.json.
"""
from __future__ import annotations

import json
import mimetypes
import os
import re
from decimal import Decimal
from pathlib import Path

import boto3
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXCEL = ROOT / "TF USA  Karwachauth-Diwali-Bhaidooj 2026  Catalouge .xlsx"
IMG_DIR = ROOT / "kdt"
CATALOG_OUT = ROOT / "scripts" / "data" / "festival-2026.json"
TABLE = os.environ.get("PRODUCTS_TABLE", "hr-ecom-products-prod")
BUCKET = os.environ["UPLOAD_BUCKET"]
CDN = os.environ["CLOUDFRONT_DOMAIN"].replace("https://", "").replace("http://", "").strip("/")
DRY_RUN = os.environ.get("DRY_RUN", "") in {"1", "true"}
PAUSE_TAG = "rakhi-season-paused-2026-10"
OLD_OOS_TAG = "usarakhi-oos-2026-08-24"
REGION = os.environ.get("AWS_DEFAULT_REGION", "us-east-1")

NAMES = {
    "TFUSBD2601": ("Bhai Dooj Kaju Katli & Kumkum Gift Box", "bhai-dooj"),
    "TFUSBD2602": ("Bhai Dooj Almonds & Pistachios Gift Box", "bhai-dooj"),
    "TFUSBD2603": ("Bhai Dooj Kaju Katli & Ferrero Rocher Box", "bhai-dooj"),
    "TFUSBD2604": ("Bhai Dooj Kaju Katli, Almonds & Kumkum Hamper", "bhai-dooj"),
    "TFUSBD2605": ("Bhai Dooj Kaju Katli & Cashew Gift Box", "bhai-dooj"),
    "TFUSBD2606": ("Bhai Dooj Mixed Dry Fruits Gift Box (266 g)", "bhai-dooj"),
    "TFUSBD2607": ("Bhai Dooj Mixed Dry Fruits Gift Box (131 g)", "bhai-dooj"),
    "TFUSBD2608": ("Bhai Dooj Ferrero, Almonds & Kumkum Box", "bhai-dooj"),
    "TFUSBD2609": ("Bhai Dooj Pistachio Gift Box", "bhai-dooj"),
    "TFUSBD2610": ("Bhai Dooj Puja Thali & Pistachio Set", "bhai-dooj"),
    "TFUSBD2611": ("Bhai Dooj Skittles Chocolate Gift", "bhai-dooj"),
    "TFUSBD2612": ("Bhai Dooj Skittles & Mixed Nuts Box", "bhai-dooj"),
    "TFUSBD2613": ("Bhai Dooj M&M's & Ferrero Rocher Box", "bhai-dooj"),
    "TFUSBD2614": ("Bhai Dooj M&M's Chocolate Gift", "bhai-dooj"),
    "TFUSBD2615": ("Bhai Dooj Pooja Set with Kaju Katli & Cashews", "bhai-dooj"),
    "TFUSBD2616": ("Bhai Dooj Pooja Set with Besan Laddoo", "bhai-dooj"),
    "TFUSBD2617": ("Bhai Dooj M&M's & Markers Kids Gift", "bhai-dooj"),
    "TFUSBD2618": ("Bhai Dooj Mixed Nuts Gift Box", "bhai-dooj"),
    "TFUSBD2619": ("Bhai Dooj M&M's & Skittles Gift Box", "bhai-dooj"),
    "TFUSDW2620": ("Diwali Toran, Rangoli & Besan Laddoo Hamper", "diwali"),
    "TFUSDW2621": ("Diwali Metal Pooja Set & Besan Laddoo Hamper", "diwali"),
    "TFUSDW2622": ("Diwali Metal Pooja Set with Kaju Katli", "diwali"),
    "TFUSDW2623": ("Diwali Shubh Labh, Rangoli & Dry Fruit Set", "diwali"),
    "TFUSDW2624": ("Diwali Pooja Set with Ferrero & Almonds", "diwali"),
    "TFUSDW2625": ("Diwali Metal Pooja Set with Kaju Katli & Pistachios", "diwali"),
}


def round_usd(amount: float) -> float:
    return int(amount * 100 + 0.5) / 100


def slugify(text: str) -> str:
    text = text.lower().replace("&", " and ")
    text = re.sub(r"[^\w\s-]", "", text)
    text = re.sub(r"[\s_-]+", "-", text).strip("-")
    return text


def canonical_sku(raw: str) -> str:
    sku = re.sub(r"\s+", "", raw.strip().upper())
    return "TFUSBD2601" if sku == "TFUSBD260_1" else sku


def clean_line(raw: str) -> str:
    line = re.sub(r"\s+", " ", raw).strip()
    line = re.sub(r"Pistachois", "Pistachios", line, flags=re.I)
    line = re.sub(r"Bhaidooj", "Bhai Dooj", line, flags=re.I)
    line = re.sub(r"Bhaidoj", "Bhai Dooj", line, flags=re.I)
    line = re.sub(r"\bkumkum\b", "Kumkum", line, flags=re.I)
    line = re.sub(r"\b(\d+)\s*gms?\b", r"\1 g", line, flags=re.I)
    line = re.sub(r"\s+", " ", line).strip()
    return (line[:1].upper() + line[1:]) if line else ""


def esc(s: str) -> str:
    return (
        s.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def meta_description(text: str, limit: int = 155) -> str:
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) <= limit:
        return text
    cut = text[: limit + 1]
    if " " in cut:
        cut = cut.rsplit(" ", 1)[0]
    return cut.rstrip(" ,;:") + "."


def images_for(sku: str) -> list[str]:
    files = []
    for path in IMG_DIR.iterdir():
        if not path.is_file():
            continue
        stem = path.stem
        base = stem.rsplit("_", 1)[0] if "_" in stem else stem
        if base.upper() == sku:
            files.append(path.name)
    def sort_key(name: str) -> tuple[int, str]:
        match = re.search(r"_(\d+)\.[^.]+$", name)
        return (int(match.group(1)) if match else 0, name)
    return sorted(files, key=sort_key)


def load_sheet() -> list[dict]:
    wb = load_workbook(EXCEL, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    products = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        data = {headers[i]: row[i] for i in range(len(headers)) if headers[i]}
        sku_raw = str(data.get("sku code") or "").strip()
        if not sku_raw:
            continue
        sku = canonical_sku(sku_raw)
        tp = float(data["tp"])
        items = []
        for key, value in data.items():
            if key.startswith("product details") and value and str(value).strip():
                cleaned = clean_line(str(value))
                if cleaned and cleaned not in items:
                    items.append(cleaned)
        name, category = NAMES[sku]
        products.append(
            {
                "sku": sku,
                "vendorCost": tp,
                "items": items,
                "name": name,
                "slug": slugify(name),
                "categorySlug": category,
            }
        )
    return products


def description(name: str, category: str, items: list[str], sku: str) -> str:
    festival = "Diwali" if category == "diwali" else "Bhai Dooj"
    lis = "".join(f"<li>{esc(item)}</li>" for item in items)
    return "\n".join(
        [
            f"<p><strong>{esc(name)}</strong> is a ready-to-gift <strong>{festival}</strong> hamper for delivery across the United States. Family overseas can send this box to a US address — it ships domestically, so it does not wait on international customs.</p>",
            "<p><strong>What's included:</strong></p>",
            f"<ul>{lis}</ul>",
            "<p>Everything above is packed together. Weights and piece counts are as listed.</p>",
            f"<p>SKU: {esc(sku)}</p>",
        ]
    )


def tags_for(items: list[str], category: str) -> list[str]:
    blob = " ".join(items).lower()
    tags = ["festival-2026", "gift-hamper", "usa-delivery", category, "orange-county"]
    if re.search(r"kaju|laddoo|ladoo|mithai|sweet", blob):
        tags.append("sweets")
    if re.search(r"ferrero|chocolate|skittles|m&m", blob):
        tags.append("chocolates")
    if re.search(r"almond|pistachio|cashew|dry fruit|mixed nut", blob):
        tags.append("dry-fruits")
    if re.search(r"pooja|puja|thali|toran|rangoli|kumkum|shubh", blob):
        tags.append("pooja")
    return tags


def ddb_num(value: float) -> Decimal:
    return Decimal(str(round_usd(value)))


def main() -> None:
    from datetime import datetime, timezone

    rows = load_sheet()
    missing = [row["sku"] for row in rows if not images_for(row["sku"])]
    if missing:
        raise SystemExit(f"No images for {', '.join(missing)}")
    print(f"Festival catalog: {len(rows)} products DRY_RUN={DRY_RUN} table={TABLE}")

    session = boto3.Session(region_name=REGION)
    ddb = session.resource("dynamodb").Table(TABLE)
    s3 = session.client("s3")
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")
    keep = set()
    catalog = []

    categories = [
        (
            "bhai-dooj",
            "Bhai Dooj Gifts",
            "Bhai Dooj gift hampers for USA delivery — sweets, dry fruits, chocolates, kumkum, and greeting cards shipped domestically.",
            "Bhai Dooj Gifts USA | Sweet & Dry Fruit Hampers",
            "Shop Bhai Dooj gift boxes for USA delivery. Kaju katli, dry fruits, chocolates, and puja sets shipped within the United States.",
            1,
        ),
        (
            "diwali",
            "Diwali Gifts",
            "Diwali gift hampers for USA delivery — pooja sets, toran, rangoli, sweets, and dry fruits shipped domestically.",
            "Diwali Gifts USA | Pooja Sets & Festive Hampers",
            "Shop Diwali hampers for USA delivery. Metal pooja sets, toran, rangoli, besan laddoo, and dry fruits shipped within the United States.",
            2,
        ),
    ]
    for slug, name, desc, seo_title, seo_desc, sort_order in categories:
        item = {
            "name": name,
            "slug": slug,
            "description": desc,
            "seoTitle": seo_title,
            "seoDescription": meta_description(seo_desc),
            "published": True,
            "sortOrder": sort_order,
            "PK": f"CATEGORY#{slug}",
            "SK": "META",
            "GSI1PK": "ENTITY#CATEGORY",
            "GSI1SK": f"{sort_order:06d}#{slug}",
            "createdAt": ts,
            "updatedAt": ts,
        }
        print(f"category {slug}")
        if not DRY_RUN:
            existing = ddb.get_item(Key={"PK": item["PK"], "SK": "META"}).get("Item")
            if existing and existing.get("createdAt"):
                item["createdAt"] = existing["createdAt"]
            ddb.put_item(Item=item)

    for row in rows:
        cost = row["vendorCost"]
        price = round_usd(cost * 1.35)
        compare = round_usd(cost * 1.75)
        files = images_for(row["sku"])
        urls = []
        for file in files:
            key = f"uploads/orange-county/{row['sku']}/{file}"
            urls.append(f"https://{CDN}/{key}")
            if DRY_RUN:
                continue
            path = IMG_DIR / file
            content_type = mimetypes.guess_type(file)[0] or "image/jpeg"
            s3.upload_file(
                str(path),
                BUCKET,
                key,
                ExtraArgs={"ContentType": content_type, "CacheControl": "public, max-age=31536000, immutable"},
            )
        heavy = row["categorySlug"] == "diwali"
        item = {
            "name": row["name"],
            "slug": row["slug"],
            "description": description(row["name"], row["categorySlug"], row["items"], row["sku"]),
            "price": price,
            "compareAtPrice": compare,
            "currency": "USD",
            "categorySlug": row["categorySlug"],
            "images": urls,
            "sku": row["sku"],
            "inventory": 500,
            "tags": tags_for(row["items"], row["categorySlug"]),
            "vendorSlug": "orange-county",
            "vendorCost": cost,
            "seoTitle": f"{row['name']} | USA Delivery",
            "seoDescription": meta_description(
                f"{row['name']} for USA delivery. Includes {', '.join(row['items'][:3])}. Ships domestically within the United States."
            ),
            "published": True,
            "weightOz": 48 if heavy else 32,
            "lengthIn": 12 if heavy else 10,
            "widthIn": 10 if heavy else 8,
            "heightIn": 4,
            "PK": f"PRODUCT#{row['slug']}",
            "SK": "META",
            "GSI1PK": f"CATEGORY#{row['categorySlug']}",
            "GSI1SK": f"PRODUCT#{row['slug']}",
            "createdAt": ts,
            "updatedAt": ts,
        }
        keep.add(row["slug"])
        catalog.append(item)
        print(f"• {row['sku']} cost={cost} sale={price} list={compare} images={len(urls)} {row['slug']}")
        if DRY_RUN:
            continue
        existing = ddb.get_item(Key={"PK": item["PK"], "SK": "META"}).get("Item")
        db_item = dict(item)
        for key in ("price", "compareAtPrice", "vendorCost", "weightOz", "lengthIn", "widthIn", "heightIn"):
            db_item[key] = ddb_num(float(db_item[key]))
        db_item["inventory"] = 500
        if existing:
            if existing.get("createdAt"):
                db_item["createdAt"] = existing["createdAt"]
            if existing.get("unitsSold") is not None:
                db_item["unitsSold"] = existing["unitsSold"]
            previous = [str(src) for src in existing.get("images") or []]
            merged = []
            seen = set()
            for src in urls + previous:
                if src not in seen:
                    seen.add(src)
                    merged.append(src)
            db_item["images"] = merged
        ddb.put_item(Item=db_item)

    CATALOG_OUT.write_text(json.dumps({"generatedAt": ts, "products": catalog}, indent=2))
    print(f"Wrote {CATALOG_OUT}")

    paused = 0
    unpublished = 0
    start_key = None
    while True:
        kwargs = {}
        if start_key:
            kwargs["ExclusiveStartKey"] = start_key
        page = ddb.scan(**kwargs)
        for item in page.get("Items", []):
            pk = str(item.get("PK", ""))
            sk = str(item.get("SK", ""))
            if sk != "META":
                continue
            slug = item.get("slug")
            if pk.startswith("CATEGORY#"):
                if not slug or slug in {"bhai-dooj", "diwali"} or item.get("published") is False:
                    continue
                print(f"{'WOULD UNPUBLISH' if DRY_RUN else 'UNPUBLISH'} CATEGORY {slug}")
                if not DRY_RUN:
                    ddb.update_item(
                        Key={"PK": pk, "SK": sk},
                        UpdateExpression="SET published = :pub, updatedAt = :now",
                        ExpressionAttributeValues={":pub": False, ":now": ts},
                    )
                unpublished += 1
                continue
            if not pk.startswith("PRODUCT#") or not slug or slug in keep:
                continue
            tags = list(item.get("tags") or [])
            if int(item.get("inventory") or 0) == 0 and PAUSE_TAG in tags:
                continue
            next_tags = [tag for tag in tags if tag != OLD_OOS_TAG]
            if PAUSE_TAG not in next_tags:
                next_tags.append(PAUSE_TAG)
            print(f"{'WOULD PAUSE' if DRY_RUN else 'PAUSE'} {slug} inventory {item.get('inventory')} → 0")
            if not DRY_RUN:
                ddb.update_item(
                    Key={"PK": pk, "SK": sk},
                    UpdateExpression="SET inventory = :zero, tags = :tags, updatedAt = :now",
                    ExpressionAttributeValues={":zero": 0, ":tags": next_tags, ":now": ts},
                )
            paused += 1
        start_key = page.get("LastEvaluatedKey")
        if not start_key:
            break
    print(json.dumps({"pausedProducts": paused, "unpublishedCategories": unpublished, "kept": len(keep), "dryRun": DRY_RUN}))


if __name__ == "__main__":
    main()
